"""Generate the internship presentation from actual generated project outputs."""
import json
from pathlib import Path
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_THEME_COLOR
from data_preprocessing import PROJECT_ROOT

OUT = PROJECT_ROOT / "Student_Performance_Prediction.pptx"; FIG = PROJECT_ROOT / "reports" / "figures"
NAVY=RGBColor(15,23,42); BLUE=RGBColor(37,99,235); TEAL=RGBColor(20,184,166); SLATE=RGBColor(71,85,105); LIGHT=RGBColor(248,250,252); WHITE=RGBColor(255,255,255)
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
def box(slide,x,y,w,h,text,size=20,color=NAVY,bold=False,fill=None,align=PP_ALIGN.LEFT):
    shape=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=shape.text_frame; tf.clear(); tf.word_wrap=True
    p=tf.paragraphs[0]; p.text=text; p.alignment=align; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=color; p.font.name="Aptos"
    if fill: shape.fill.solid(); shape.fill.fore_color.rgb=fill
    return shape
def title(slide,text,subtitle=None):
    box(slide,.65,.4,12,.55,text,30,NAVY,True)
    if subtitle: box(slide,.68,1.02,11.5,.35,subtitle,14,SLATE)
def footer(slide,n): box(slide,.65,7.05,12,.22,f"Khizar Nawaz  |  Applied AI Internship  |  {n}",10,SLATE,align=PP_ALIGN.RIGHT)
def bullet(slide,x,y,text): box(slide,x,y,5.5,.5,"• "+text,18,NAVY)
def panel(slide,x,y,w,h,heading,body):
    s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); s.fill.solid(); s.fill.fore_color.rgb=LIGHT; s.line.color.rgb=RGBColor(226,232,240)
    box(slide,x+.22,y+.18,w-.44,.35,heading,18,BLUE,True); box(slide,x+.22,y+.65,w-.44,h-.8,body,14,SLATE)
# 1
s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=NAVY
box(s,.75,1.35,11.7,1.4,"AI-Based Student Performance\nPrediction System",40,WHITE,True)
box(s,.8,3.15,10,.45,"A reproducible machine-learning prototype for academic-support insights",20,RGBColor(191,219,254))
box(s,.8,5.65,5,.35,"Khizar Nawaz",18,WHITE,True); box(s,.8,6.12,5,.3,"Applied AI Internship",15,RGBColor(148,163,184)); footer(s,1)
# 2
s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"Problem Statement","Turning multiple student factors into a clear support conversation")
bullet(s,.8,1.7,"Academic performance is influenced by learning, support, attendance, and social context."); bullet(s,.8,2.45,"Manual review may not combine many signals consistently."); bullet(s,.8,3.2,"Early estimates can help educators discuss possible additional support."); bullet(s,.8,3.95,"Machine learning can identify historical patterns — not make definitive judgments.")
panel(s,8.1,1.85,4.1,3.65,"Project boundary","Educational prototype\n\nNot a clinical or psychological diagnosis.\n\nDoes not replace teachers or guarantee outcomes."); footer(s,2)
# 3
s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"Proposed Solution","An end-to-end local machine-learning workflow")
steps=["Student\nData","Preprocess","Regression\nModels","Evaluate","Best\nModel","Streamlit\nPrediction"]
for i,step in enumerate(steps):
    x=.55+i*2.1; sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(2.8),Inches(1.55),Inches(.95)); sh.fill.solid(); sh.fill.fore_color.rgb=BLUE if i in (0,5) else LIGHT; sh.line.color.rgb=BLUE
    box(s,x+.08,3.04,1.4,.45,step,16,WHITE if i in (0,5) else NAVY,True,align=PP_ALIGN.CENTER)
    if i<5: box(s,x+1.6,3.05,.35,.3,"→",22,TEAL,True,align=PP_ALIGN.CENTER)
box(s,1.2,4.65,10.8,.55,"Output: predicted final grade / 20  •  performance category  •  academic-support indicator",19,SLATE,align=PP_ALIGN.CENTER); footer(s,3)
# 4
s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"Dataset & Features","Official UCI Student Performance Dataset — Mathematics subset")
panel(s,.8,1.55,3.6,3.9,"Dataset facts","395 student records\n33 columns\nTarget: G3 final grade\nGrade scale: 0–20")
panel(s,4.85,1.55,3.6,3.9,"Feature groups","Demographics\nFamily & school support\nStudy behaviour\nAttendance\nPrior-term grades")
panel(s,8.9,1.55,3.6,3.9,"Timing decision","G1 and G2 are valid when prior-term grades exist.\n\nA second scenario excludes them for earlier prediction.")
box(s,.8,5.9,11.7,.35,"Source: UCI Machine Learning Repository, Student Performance dataset (dataset id 320).",13,SLATE); footer(s,4)
# 5
s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"Model Comparison & Results","Actual metrics from the fixed 80/20 holdout split (random_state = 42)")
results=pd.read_csv(PROJECT_ROOT/'reports'/'model_comparison.csv'); rows,cols=results.shape
x,y,w,h=.8,1.55,7.0,3.0; table=s.shapes.add_table(rows+1,cols,Inches(x),Inches(y),Inches(w),Inches(h)).table
headers=list(results.columns)
for j,header in enumerate(headers): cell=table.cell(0,j); cell.text=header; cell.fill.solid(); cell.fill.fore_color.rgb=BLUE
for i,row in results.iterrows():
    for j,header in enumerate(headers): table.cell(i+1,j).text=str(row[header]) if header=='Model' else f"{row[header]:.3f}"
for row in table.rows:
    for cell in row.cells:
        for p in cell.text_frame.paragraphs: p.font.size=Pt(13); p.font.name="Aptos"; p.font.color.rgb=NAVY
best=results.iloc[0]; panel(s,8.4,1.65,3.9,2.65,"Selected model",f"{best['Model']}\n\nMAE: {best['MAE']:.3f}\nRMSE: {best['RMSE']:.3f}\nR²: {best['R2']:.3f}")
s.shapes.add_picture(str(FIG/'model_comparison.png'), Inches(8.25), Inches(4.45), width=Inches(4.2)); footer(s,5)
# 6
s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"End-to-End Workflow","Reproducible and leakage-safe pipeline")
items=["Load UCI data","Check schema, missing values & duplicates","Split train/test data","Encode categories and scale numbers in pipeline","Train and compare models","Save selected full pipeline","Evaluate, explain and deploy"]
for i,item in enumerate(items):
    y=1.35+i*.68; box(s,1.15,y,.45,.4,str(i+1),18,WHITE,True,BLUE,PP_ALIGN.CENTER); box(s,1.8,y,8.4,.4,item,18,NAVY);
    if i<6: box(s,1.25,y+.39,.25,.25,"↓",16,TEAL,True,align=PP_ALIGN.CENTER)
s.shapes.add_picture(str(FIG/'feature_importance.png'), Inches(9.35), Inches(1.45), width=Inches(3.55)); footer(s,6)
# 7
s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"Working Demo","What the Streamlit application demonstrates")
for i,text in enumerate(["Enter grouped student information", "Click Predict Performance", "Show predicted final grade", "Show performance category", "Show academic-support indicator", "Explain important model features"]):
    panel(s,.8+(i%2)*6.15,1.5+(i//2)*1.55,5.5,1.05,f"{i+1}",text)
box(s,.85,6.25,11.5,.35,"Live demonstration is intentionally used instead of a fabricated screenshot.",14,SLATE,align=PP_ALIGN.CENTER); footer(s,7)
# 8
s=prs.slides.add_slide(prs.slide_layouts[6]); title(s,"Conclusion & Key Learnings","A complete Applied AI workflow, ready for local demonstration")
bullet(s,.85,1.6,"Compared multiple regression models using standard, actual evaluation metrics."); bullet(s,.85,2.35,"Integrated the best model into a clear Streamlit application."); bullet(s,.85,3.1,"Made model drivers visible through feature-importance analysis."); bullet(s,.85,3.85,"Documented responsible use: support discussions, not definitive decisions.")
panel(s,8.25,1.75,4.05,2.75,"Future improvements","More representative data\nFairness assessment\nPrediction uncertainty\nEducator feedback loop")
box(s,.85,5.75,11.5,.55,"Thank you",30,BLUE,True,align=PP_ALIGN.CENTER); footer(s,8)
prs.save(OUT); print(OUT)
